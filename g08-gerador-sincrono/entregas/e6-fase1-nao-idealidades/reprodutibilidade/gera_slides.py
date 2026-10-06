# -*- coding: utf-8 -*-
"""
ECOM060 - Grupo 08 - Fase 1, Entrega 02
Gera a apresentacao (25 min) em ../apresentacao/Equipe_08_F1E02_apresentacao.pptx.
Abre no PowerPoint e importa no Google Slides (File > Import slides).

Usa as figuras PNG de ../relatorio/figuras e os numeros de resultados_camadas.json
(e de resultados_typhoon.json, se existir). Rode de novo depois do TyphoonSim: as figuras
e os valores "pendente" sao trocados automaticamente pelos da simulacao.

Uso:  python gera_slides.py        (precisa de python-pptx)
"""
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt

import parametros as P

AQUI = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(AQUI, '..', 'relatorio', 'figuras')
SAIDA = os.path.join(AQUI, '..', 'apresentacao', 'Equipe_08_F1E02_apresentacao.pptx')
R = json.load(open(os.path.join(AQUI, 'resultados_camadas.json'), encoding='utf-8'))
fty = os.path.join(AQUI, 'resultados_typhoon.json')
TY = json.load(open(fty, encoding='utf-8')) if os.path.exists(fty) else {}

AZUL, LARANJA, AQUA = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A)
TINTA, TINTA2, CINZA = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x8A, 0x89, 0x83)
CLARO = RGBColor(0xF3, 0xF2, 0xEE)
W, H = Cm(33.867), Cm(19.05)          # 16:9
RODAPE = 'ECOM060 · Grupo 08 · Fase 1 — Entrega 02'


def br(v, nd):
    return ('%.*f' % (nd, v)).replace('.', ',')


def cient(v):
    m, e = ('%.1e' % v).split('e')
    return '%s×10%s' % (m.replace('.', ','), ''.join('⁰¹²³⁴⁵⁶⁷⁸⁹⁻'['0123456789-'.index(c)] for c in str(int(e))))


def ty(caso, chave, nd):
    try:
        return br(TY[caso][chave], nd)
    except KeyError:
        return 'pendente'


def fig(nome_typhoon, nome_python):
    f = os.path.join(FIG, nome_typhoon + '.png')
    return (f, True) if os.path.exists(f) else (os.path.join(FIG, nome_python + '.png'), False)


prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BRANCO = prs.slide_layouts[6]
n_slide = [0]


def texto(sl, x, y, w, h, linhas, tam=16, cor=TINTA, negrito=False, alinh=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP):
    """linhas: str ou lista de str / (str, dict) com chaves tam, cor, negrito, marcador."""
    tb = sl.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Cm(0.1); tf.margin_top = tf.margin_bottom = Cm(0.05)
    if isinstance(linhas, str):
        linhas = [linhas]
    for i, ln in enumerate(linhas):
        opt = {}
        if isinstance(ln, tuple):
            ln, opt = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinh
        p.space_after = Pt(opt.get('depois', 6))
        r = p.add_run()
        r.text = ('•  ' if opt.get('marcador') else '') + ln
        r.font.size = Pt(opt.get('tam', tam)); r.font.bold = opt.get('negrito', negrito)
        r.font.color.rgb = opt.get('cor', cor); r.font.name = 'Calibri'
    return tb


def caixa(sl, x, y, w, h, cor_fundo=CLARO, cor_borda=None):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.08
    s.fill.solid(); s.fill.fore_color.rgb = cor_fundo
    if cor_borda is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = cor_borda; s.line.width = Pt(1.2)
    s.shadow.inherit = False
    return s


def novo_slide(tag, titulo, sub=None):
    n_slide[0] += 1
    sl = prs.slides.add_slide(BRANCO)
    if tag:
        t = caixa(sl, Cm(1.2), Cm(0.9), Cm(1.9), Cm(1.1), cor_fundo=AZUL)
        texto(sl, Cm(1.2), Cm(0.9), Cm(1.9), Cm(1.1), tag, tam=16, cor=RGBColor(255, 255, 255),
              negrito=True, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    texto(sl, Cm(3.5 if tag else 1.2), Cm(0.75), Cm(29), Cm(1.4), titulo, tam=28, negrito=True,
          ancora=MSO_ANCHOR.MIDDLE)
    if sub:
        texto(sl, Cm(3.5 if tag else 1.2), Cm(2.15), Cm(29), Cm(0.9), sub, tam=14, cor=TINTA2)
    texto(sl, Cm(1.2), Cm(18.0), Cm(25), Cm(0.7), RODAPE, tam=10, cor=CINZA)
    texto(sl, Cm(30.5), Cm(18.0), Cm(2.2), Cm(0.7), str(n_slide[0]), tam=10, cor=CINZA, alinh=PP_ALIGN.RIGHT)
    return sl


def imagem(sl, caminho, x, y, w=None, h=None):
    return sl.shapes.add_picture(caminho, x, y, width=w, height=h)


def numero(sl, x, y, w, valor, legenda, cor=AZUL):
    texto(sl, x, y, w, Cm(1.6), valor, tam=32, negrito=True, cor=cor)
    texto(sl, x, y + Cm(1.55), w, Cm(1.3), legenda, tam=13, cor=TINTA2)


def seta(sl, x1, y1, x2, y2, cor=CINZA):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = cor; c.line.width = Pt(1.8)
    c.line._get_or_add_ln().append(_ponta())
    return c


def _ponta():
    from pptx.oxml.ns import qn
    from lxml import etree
    e = etree.SubElement(etree.Element('dummy'), qn('a:tailEnd'))
    e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    return e


C, F, R6, R7, R8 = R['casos']['completa'], R['casos_grade_fina'], R['r6'], R['r7'], R['r8']
SB = R['r8_circuito_sem_buffer']; D = R['desvio_contra_E01']

# ============================================================== 1. capa
n_slide[0] += 1
sl = prs.slides.add_slide(BRANCO)
texto(sl, Cm(2), Cm(2.2), Cm(30), Cm(0.8), 'ECOM060 · INSTRUMENTAÇÃO ELETRÔNICA · FASE 1', tam=14, cor=AZUL, negrito=True)
texto(sl, Cm(2), Cm(4.6), Cm(30), Cm(2.2), 'Entrega 02 — Não idealidades', tam=44, negrito=True)
texto(sl, Cm(2), Cm(6.9), Cm(30), Cm(1.4), 'Potenciômetro de precisão + laço de corrente 4–20 mA, agora com o sensor real',
      tam=20, cor=TINTA2)
# faixa com as tres camadas
for i, (tag, nome, cor) in enumerate([('R6', 'estática: resolução + histerese do contato', AZUL),
                                       ('R7', 'dinâmica: ordem zero', AQUA),
                                       ('R8', 'espectral: ruído de contato 1/f', LARANJA)]):
    caixa(sl, Cm(2 + i * 10.1), Cm(9.6), Cm(9.6), Cm(2.0), cor_borda=cor)
    texto(sl, Cm(2.3 + i * 10.1), Cm(9.75), Cm(9.1), Cm(1.8), [(tag, dict(tam=16, negrito=True, cor=cor, depois=0)),
                                                               (nome, dict(tam=13, cor=TINTA))])
texto(sl, Cm(2), Cm(13.4), Cm(30), Cm(2.4),
      [('Grupo 08', dict(tam=18, negrito=True)),
       ('Cleiber de Meireles da Silva Junnior · Leonardo Lima Barbosa Pereira · Willian Tcheldon Oliveira Santos', dict(tam=14)),
       ('Prof. Maurício Beltrão Rossiter · UFAL · 2026.2', dict(tam=14, cor=TINTA2))])

# ============================================================== 2. o que muda
sl = novo_slide(None, 'Da cadeia ideal para o sensor real', 'O circuito da Entrega 01 não muda — as camadas agem sobre o potenciômetro (R10)')
numero(sl, Cm(1.6), Cm(3.6), Cm(9.5), '0,120 % FE', 'erro total da cadeia ideal (E01):\ntodo ele dos resistores comerciais')
numero(sl, Cm(12.0), Cm(3.6), Cm(9.5), '3 camadas', 'atribuídas ao G08\n(slide 15 da apresentação da fase)', cor=LARANJA)
numero(sl, Cm(22.4), Cm(3.6), Cm(10), '< 1 % FE', 'a especificação da família C\ncontinua valendo', cor=AQUA)
# diagrama da cadeia com as camadas
blocos = [('eixo θ', CINZA), ('folga\n1,0°', AZUL), ('resolução\n0,72°', AZUL), ('ordem zero\nK·x', AQUA),
          ('+ ruído de contato\nΔRc · I_cursor', LARANJA), ('circuito da E01\nOPA333 → XTR115', TINTA2), ('4–20 mA', CINZA)]
x0, y0, bw, bh, gap = Cm(1.4), Cm(9.6), Cm(4.0), Cm(2.2), Cm(0.65)
for i, (nm, cor) in enumerate(blocos):
    x = x0 + i * (bw + gap)
    caixa(sl, x, y0, bw, bh, cor_borda=cor)
    texto(sl, x, y0, bw, bh, nm, tam=13, cor=TINTA, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
    if i < len(blocos) - 1:
        seta(sl, x + bw, y0 + bh // 2, x + bw + gap, y0 + bh // 2)
texto(sl, Cm(1.4), Cm(12.6), Cm(31), Cm(4.6),
      [('Mesmo sinal de referência da E01: varredura 0 → 3600° → 0 em 20 s (360 °/s, 60 rpm).', dict(marcador=True)),
       ('Todo valor vem de datasheet citado: Bourns 3590 (resolução 0,020 %, backlash 1,0°, ENR 100 Ω), TI OPA333 e XTR115.', dict(marcador=True)),
       ('Uma definição única das camadas gera o gabarito Python e o modelo do TyphoonSim — conferidos amostra a amostra.', dict(marcador=True))],
      tam=15)

# ============================================================== 3. R6 resolucao
sl = novo_slide('R6', 'Resolução: o cursor anda de espira em espira', 'Cap. 3 — fio enrolado com 0,020 % de resolução (Bourns 3590, 10 kΩ)')
f = os.path.join(FIG, 'fig-07-r6-estatica.png')
imagem(sl, f, Cm(1.2), Cm(3.5), w=Cm(19.5))
numero(sl, Cm(22.0), Cm(3.6), Cm(11), '%d espiras' % R6['N_passos'], '1 / 0,0002 = posições discretas do cursor')
numero(sl, Cm(22.0), Cm(7.0), Cm(11), '%s°' % br(R6['passo_graus'], 2), 'por degrau = %s mV no cursor\n= %s µA no laço' % (br(R6['passo_mV'], 1), br(R6['passo_uA'], 1)))
numero(sl, Cm(22.0), Cm(10.9), Cm(11), '±%s°' % br(R6['erro_quantizacao_max_graus'], 2), 'erro máximo = ±%s %% FE\n(meio degrau)' % br(R6['erro_quantizacao_max_pcFS'], 3))
texto(sl, Cm(1.2), Cm(13.8), Cm(20), Cm(3.6),
      [('Modelo: x_q = (⌊x·N + ½ − φ⌋ + φ) / N', dict(negrito=True)),
       ('φ = ¼ de passo (hipótese): a 1ª espira não está exatamente em 0°. Sem isso, metade das amostras de 1 ms caía '
        'na fronteira entre espiras e o arredondamento numérico inventava uma histerese de 0,72°.', dict(tam=13, cor=TINTA2))], tam=15)

# ============================================================== 4. R6 folga + limiar
sl = novo_slide('R6', 'Histerese do contato: a folga entre eixo e cursor', 'Backlash de 1,0° máx. (Bourns 3590) — modelado como operador de folga, calibrado na subida')
caixa(sl, Cm(1.2), Cm(3.6), Cm(15.3), Cm(3.0))
texto(sl, Cm(1.5), Cm(3.75), Cm(14.8), Cm(2.8),
      [('θc[k] = max( θ, min(θ + b, θc[k−1]) )', dict(tam=18, negrito=True)),
       ('subindo: θc = θ   ·   descendo: θc = θ + b   ·   na reversão: parado', dict(tam=13, cor=TINTA2))])
numero(sl, Cm(1.6), Cm(7.2), Cm(7), '%s°' % br(F['so_folga']['histerese_graus'], 2), 'histerese só da folga')
numero(sl, Cm(9.0), Cm(7.2), Cm(7.5), '%s°' % br(F['estatica']['histerese_graus'], 2), 'folga + escada:\npula 1 ou 2 degraus', cor=LARANJA)
caixa(sl, Cm(17.6), Cm(3.6), Cm(15.0), Cm(7.0), cor_borda=AZUL)
texto(sl, Cm(18.0), Cm(3.8), Cm(14.3), Cm(6.7),
      [('Resolução não é limiar (slide 20 do professor)', dict(tam=17, negrito=True, cor=AZUL)),
       ('Resolução (VIM 4.14): o degrau de %s° — em qualquer ponto da faixa.' % br(R6['passo_graus'], 2), dict(marcador=True)),
       ('Limiar (VIM 4.16) aparece na reversão: o eixo precisa vencer a folga e o resto da espira para a leitura mudar.', dict(marcador=True)),
       ('Zona morta: de 1,00° a %s° (na nossa reversão: %s°).' % (br(R6['zona_morta_max_graus'], 2), br(F['estatica']['zona_morta_reversao_graus'], 2)), dict(marcador=True)),
       ('A folga cria um limiar sem mudar a resolução.', dict(negrito=True))], tam=14)
texto(sl, Cm(1.2), Cm(11.6), Cm(31.5), Cm(5.8),
      [('Por que 1,44° e não 1°?', dict(negrito=True, tam=16)),
       ('Na descida, a escada quantiza θ + 1°. Como 1° fica entre 1 degrau (0,72°) e 2 degraus (1,44°), a diferença '
        'entre subida e descida só pode valer um ou dois degraus ao longo da faixa.', dict(tam=14, cor=TINTA2)),
       ('Desvio total da camada estática: ±0,36° na subida, entre +0,64° e +%s° na descida.' % br(D['estatica']['max_graus'], 2), dict(tam=14))])

# ============================================================== 5. R7 ordem zero
sl = novo_slide('R7', 'Dinâmica: ordem zero, com números', 'Cap. 4 — o divisor v = x·Vs é algébrico: nenhum elemento armazena energia entre ângulo e tensão')
imagem(sl, os.path.join(FIG, 'fig-08-r7-dinamica.png'), Cm(1.2), Cm(3.5), w=Cm(18.5))
texto(sl, Cm(20.8), Cm(3.5), Cm(12), Cm(6.6),
      [('Polos reais da cadeia (datasheets)', dict(negrito=True, tam=16)),
       ('entrada do buffer: %s ns' % br(R7['tau_entrada_ns'], 1), dict(marcador=True)),
       ('OPA333 (GBW 350 kHz): %s µs' % br(R7['tau_buffer_us'], 3), dict(marcador=True)),
       ('XTR115 (380 kHz): %s µs' % br(R7['tau_xtr_us'], 3), dict(marcador=True)),
       ('total: τ = %s µs → 99 %% em %s µs' % (br(R7['tau_total_us'], 2), br(R7['t_99_us'], 2)), dict(negrito=True))], tam=15)
numero(sl, Cm(20.8), Cm(10.6), Cm(12), cient(R7['atraso_rampa_graus']) + '°', 'atraso de rampa τ·dθ/dt na varredura', cor=AQUA)
texto(sl, Cm(1.2), Cm(13.4), Cm(31.5), Cm(4.2),
      [('No TyphoonSim a camada é só o ganho K = 0,694 mV/grau (sensibilidade do sensor ≠ ganho da cadeia de 4,44 µA/grau).', dict(marcador=True)),
       ('O passo de 1 ms é ~%d× maior que τ: um polo de µs seria invisível.' % round(R7['razao_passo_typhoon_sobre_tau'], -1), dict(marcador=True)),
       ('Cuidado: o atraso NUMÉRICO de 2 ms do simulador vale 0,72° a 360 °/s — exatamente um degrau. Compensado na análise.', dict(marcador=True, negrito=True))],
      tam=14)

# ============================================================== 6. R8 o que e o ENR
sl = novo_slide('R8', 'Ruído de contato: o que o datasheet realmente diz', 'Cap. 5 — 100 Ω ENR máximo (Bourns 3590)')
caixa(sl, Cm(1.2), Cm(3.6), Cm(14.6), Cm(6.4), cor_borda=LARANJA)
texto(sl, Cm(1.6), Cm(3.8), Cm(13.9), Cm(6.1),
      [('ENR = Equivalent Noise Resistance', dict(negrito=True, tam=17, cor=LARANJA)),
       ('resistência parasita transitória entre cursor e trilha, só com o eixo girando', dict(marcador=True)),
       ('ensaiada com 1 mA passando pelo cursor (catálogo Bourns, p. 150)', dict(marcador=True)),
       ('É uma RESISTÊNCIA: só vira tensão com corrente no contato', dict(negrito=True))], tam=14)
caixa(sl, Cm(1.2), Cm(10.6), Cm(14.6), Cm(2.0))
texto(sl, Cm(1.2), Cm(10.6), Cm(14.6), Cm(2.0), 'v_n = ΔRc(t) · I_cursor', tam=24, negrito=True,
      alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
texto(sl, Cm(1.2), Cm(13.2), Cm(14.6), Cm(4.4),
      [('Modelo: gaussiano → filtro 1/f (Kellet). No Typhoon: Random Source + 6 funções de transferência nativas.', dict(marcador=True)),
       ('Inclinação medida: %s déc./déc. · pico 100 Ω · eficaz %s Ω.' % (br(R8['inclinacao_psd_dec'], 2), br(R8['dRc_rms_ohm'], 1)), dict(marcador=True)),
       ('Só com o cursor em movimento (na zona morta da folga, não há ruído).', dict(marcador=True))], tam=14)
imagem(sl, os.path.join(FIG, 'fig-09-r8-espectral.png'), Cm(16.6), Cm(3.5), w=Cm(16.2))

# ============================================================== 7. R8 resultado (buffer)
sl = novo_slide('R8', 'O buffer da E01 também mata o ruído de contato', 'Mesmo ENR, três correntes no cursor')
numero(sl, Cm(1.6), Cm(3.8), Cm(10), '%s nV' % br(R8['v_ruido_buffer_nV'], 0), 'com buffer (I_B ≤ 200 pA)\n= %s %% FE' % cient(R8['erro_buffer_pcFS']), cor=AZUL)
numero(sl, Cm(12.0), Cm(3.8), Cm(10), '%s %% FE' % br(R8['erro_sem_buffer_pcFS'], 2), 'sem buffer: cursor alimenta R_IN\n(64 % do limite, só de ruído)', cor=LARANJA)
numero(sl, Cm(22.4), Cm(3.8), Cm(10), '%s %% FE' % br(R8['erro_ensaio_pcFS'], 1), 'nas condições do ensaio\ndo datasheet (1 mA)', cor=CINZA)
caixa(sl, Cm(1.2), Cm(8.4), Cm(31.5), Cm(3.0), cor_fundo=RGBColor(0xE8, 0xF1, 0xFB))
texto(sl, Cm(1.6), Cm(8.5), Cm(30.8), Cm(2.8),
      [('A mesma causa — corrente passando pelo cursor — está por trás de duas não idealidades:', dict(tam=16)),
       ('efeito de carga (8,3 % FE, E01) e ruído de contato (0,64 % FE, E02). Um único buffer resolve as duas.', dict(tam=16, negrito=True))])
texto(sl, Cm(1.2), Cm(12.2), Cm(31.5), Cm(5),
      [('Pergunta provável: “então o ruído não aparece?”', dict(negrito=True)),
       ('Ele foi modelado e aplicado nos dois simuladores. Não aparece porque o projeto da E01 já o suprimiu — '
        'e o próximo slide mostra isso acontecendo no próprio circuito.', dict(cor=TINTA2))], tam=15)

# ============================================================== 8. R8 contrafactual no circuito
f8, eh_ty8 = fig('fig-14-r8-typhoon', 'fig-13-r8-sem-buffer')
sl = novo_slide('R8', 'Prova no circuito: potenciômetro sem buffer',
                'Trilha e contato como resistores variáveis do TyphoonSim; cursor ligado direto em R_IN')
im = imagem(sl, f8, Cm(1.2), Cm(3.5), w=Cm(21.5))
texto(sl, Cm(1.2), im.top + im.height + Cm(0.1), Cm(21.5), Cm(0.8), 'fonte: ' + ('TyphoonSim' if eh_ty8 else 'gabarito Python (troca pelo TyphoonSim ao regerar)'),
      tam=11, cor=CINZA)
numero(sl, Cm(23.6), Cm(3.6), Cm(9.5), '%s %% FE' % br(SB['erro_lin_subida_pcFS'], 2), 'efeito de carga\n(TyphoonSim: %s)' % ty('r8_sem_buffer', 'erro_lin_subida_pcFS', 2), cor=LARANJA)
numero(sl, Cm(23.6), Cm(7.4), Cm(9.5), '%s %% FE' % br(SB['ruido_pico_pcFS'], 2), 'pico do ruído de contato\n(TyphoonSim: %s)' % ty('r8_sem_buffer', 'ruido_pico_pcFS', 2), cor=LARANJA)
texto(sl, Cm(23.6), Cm(11.4), Cm(9.5), Cm(6),
      [('O ruído cresce com x: a corrente no cursor é ∝ x.', dict(marcador=True)),
       ('R_mín = 10 Ω nas pontas (datasheet); contato de base 110 Ω (hipótese, só para nunca zerar).', dict(marcador=True))], tam=13)
texto(sl, Cm(1.2), im.top + im.height + Cm(1.0), Cm(21.5), Cm(4),
      [('Contrafactual — não é o projeto. Serve para mostrar, no simulador, o que o buffer evita.', dict(negrito=True)),
       ('Os dois efeitos saem das leis do circuito, não de uma fórmula.', dict(cor=TINTA2))], tam=15)

# ============================================================== 9. R9
sl = novo_slide('R9', 'Comparação por camada contra o sinal ideal da E01', 'Mesma varredura, mesma base de tempo — atenção às escalas')
imagem(sl, os.path.join(FIG, 'fig-10-r9-camadas.png'), Cm(1.2), Cm(3.3), h=Cm(14.4))
linhas_t = [('resolução', '±%s°' % br(D['so_resolucao']['max_graus'], 2), AZUL),
            ('folga', '%s°' % br(D['so_folga']['max_graus'], 2), AZUL),
            ('estática (as duas)', '%s°' % br(D['estatica']['max_graus'], 2), AZUL),
            ('dinâmica', cient(R7['atraso_rampa_graus']) + '°', AQUA),
            ('espectral', cient(D['espectral']['max_graus']) + '°', LARANJA)]
yy = Cm(3.8)
texto(sl, Cm(22.2), yy, Cm(10.5), Cm(0.9), 'Desvio máximo', tam=16, negrito=True)
for nm, v, cor in linhas_t:
    yy += Cm(1.25)
    texto(sl, Cm(22.2), yy, Cm(5.8), Cm(1.0), nm, tam=15, cor=TINTA2)
    texto(sl, Cm(28.0), yy, Cm(4.7), Cm(1.0), v, tam=15, negrito=True, cor=cor, alinh=PP_ALIGN.RIGHT)
texto(sl, Cm(22.2), Cm(11.4), Cm(10.5), Cm(5.5),
      [('A camada estática domina; dentro dela, a folga pesa mais que a resolução.', dict(marcador=True)),
       ('Dinâmica e ruído: 3 a 5 ordens de grandeza abaixo.', dict(marcador=True))], tam=14)

# ============================================================== 10. R10
f10, eh_ty10 = fig('fig-12-r10-typhoon', 'fig-11-r10-completa')
sl = novo_slide('R10', 'As três camadas no circuito da E01: atende', 'Critérios da Entrega 01, mesma métrica nos dois simuladores (grade de 1 ms)')
im = imagem(sl, f10, Cm(1.2), Cm(3.4), w=Cm(19.5))
texto(sl, Cm(1.2), im.top + im.height + Cm(0.1), Cm(19.5), Cm(0.8), 'figura: ' + ('TyphoonSim' if eh_ty10 else 'gabarito Python (troca pelo TyphoonSim ao regerar)'), tam=11, cor=CINZA)
cab = [('critério (< 1 % FE)', 'Python', 'TyphoonSim')]
lin = [('linearidade, subida', br(C['erro_lin_subida_pcFS'], 4), ty('completa', 'erro_lin_subida_pcFS', 4)),
       ('linearidade, banda', br(C['erro_lin_banda_pcFS'], 4), ty('completa', 'erro_lin_banda_pcFS', 4)),
       ('erro total', br(C['erro_total_pcFS'], 3), ty('completa', 'erro_total_pcFS', 3)),
       ('   (E01, ideal)', br(R['casos']['ideal']['erro_total_pcFS'], 3), ty('ideal', 'erro_total_pcFS', 3)),
       ('histerese [°]', br(C['histerese_graus'], 2), ty('completa', 'histerese_graus', 2))]
tab = sl.shapes.add_table(len(lin) + 1, 3, Cm(21.4), Cm(3.5), Cm(11.4), Cm(6.6)).table
tab.columns[0].width = Cm(5.4); tab.columns[1].width = Cm(2.8); tab.columns[2].width = Cm(3.2)
for i, row in enumerate(cab + lin):
    for j, v in enumerate(row):
        cel = tab.cell(i, j); cel.text = v
        p = cel.text_frame.paragraphs[0]; p.font.size = Pt(13); p.font.bold = (i == 0)
        p.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else TINTA
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        cel.fill.solid(); cel.fill.fore_color.rgb = AZUL if i == 0 else (CLARO if i % 2 else RGBColor(255, 255, 255))
texto(sl, Cm(21.4), Cm(10.6), Cm(11.4), Cm(1.6), 'erro total: 0,120 → %s %% FE\nmargem de %s× sobre o limite' % (br(C['erro_total_pcFS'], 3), br(100 * P.ERRO_LIN_MAX / C['erro_total_pcFS'], 1)),
      tam=15, negrito=True, cor=AQUA)
texto(sl, Cm(1.2), im.top + im.height + Cm(1.0), Cm(31.5), Cm(5),
      [('Entre as camadas, a folga é a que mais afasta o sinal do ideal (1°, só na descida).', dict(marcador=True)),
       ('O maior erro da cadeia continua sendo o ganho dos resistores comerciais da E01 (0,120 % FE).', dict(marcador=True)),
       ('A linearidade de ±0,25 % do potenciômetro (não atribuída nesta entrega) segue no orçamento da E01.', dict(marcador=True))], tam=15)

# ============================================================== 11. TyphoonSim
sl = novo_slide(None, 'Implementação no TyphoonSim', 'O que veio da E01 ficou intacto em C; tudo o que é da E02 é bloco nativo')
for i, (tit, cor, itens) in enumerate([
        ('E01 — C function', CINZA, ['referência (varredura) e receptor',
                                     'potenciômetro: x = θ/3600, v = x·V(VREF)',
                                     'circuito: buffer, R0, R_IN, XTR115, laço']),
        ('E02 — blocos nativos', AZUL, ['folga: Sum + Min Max + Unit Delay',
                                        'resolução: Gain + Sum + Round (floor)',
                                        'ruído: Random Source + 6 Discrete Transfer Function (1/f)',
                                        'movimento: Abs + Sign · liga/desliga por Constant']),
        ('Conferência', AQUA, ['blocos emulados + C da E01 (gcc) = Python (≤ 4×10⁻¹⁶ V)',
                               'Random Source sem semente: ruído comparado pela estatística',
                               'filtro 1/f nativo conferido a partir do branco sorteado'])]):
    x = Cm(1.2 + i * 10.6)
    caixa(sl, x, Cm(3.8), Cm(10.0), Cm(8.6), cor_borda=cor)
    texto(sl, x + Cm(0.4), Cm(4.0), Cm(9.3), Cm(8.3), [(tit, dict(tam=17, negrito=True, cor=cor))] +
          [(t, dict(marcador=True, tam=14)) for t in itens])
texto(sl, Cm(1.2), Cm(13.2), Cm(31.5), Cm(4.3),
      [('Atraso numérico: 2 amostras (2 ms) entre θ e I_LOOP, como na E01 — vale um degrau de resolução a 360 °/s, '
        'por isso a análise realinha antes de comparar.', dict(marcador=True)),
       ('Resultado do TyphoonSim: %s.' % ('parte determinística: diferença contra o Python de %s µA' % ty('completa', 'dif_typhoon_python_max_uA', 4)
                                          if TY else 'pendente — rodar a simulação no TyphoonSim'), dict(marcador=True, negrito=True))], tam=14)

# ============================================================== 12. conclusoes
sl = novo_slide(None, 'Conclusões e próximos passos')
texto(sl, Cm(1.2), Cm(3.0), Cm(19), Cm(14),
      [('O que a Entrega 02 mostrou', dict(tam=18, negrito=True, cor=AZUL)),
       ('A cadeia real continua atendendo à família C: %s %% FE de erro total.' % br(C['erro_total_pcFS'], 3), dict(marcador=True)),
       ('A folga é a camada que mais pesa — e é ela que cria um limiar na reversão.', dict(marcador=True)),
       ('Ordem zero confirmada com os polos reais: τ = %s µs.' % br(R7['tau_total_us'], 2), dict(marcador=True)),
       ('O buffer da E01 neutraliza o ruído de contato; sem ele: %s %% FE de ruído + %s %% FE de carga.'
        % (br(SB['ruido_pico_pcFS'], 2), br(SB['erro_lin_subida_pcFS'], 1)), dict(marcador=True))], tam=16)
caixa(sl, Cm(21.0), Cm(3.0), Cm(11.7), Cm(10.6), cor_borda=LARANJA)
texto(sl, Cm(21.4), Cm(3.2), Cm(11.0), Cm(10.2),
      [('Entrega 03 (19–21/out)', dict(tam=18, negrito=True, cor=LARANJA)),
       ('degrau, condição inicial e transitório sobre esta cadeia', dict(marcador=True)),
       ('ordem zero → os três ensaios devem sair parecidos (esperado, não falha)', dict(marcador=True)),
       ('o que deve aparecer: a zona morta nas inversões e o atraso numérico de 2 ms nos degraus', dict(marcador=True))], tam=15)
texto(sl, Cm(1.2), Cm(15.6), Cm(31.5), Cm(1.6), 'Todos os números vêm de datasheet citado e dos scripts em reprodutibilidade/.',
      tam=13, cor=TINTA2)

# ============================================================== 13. apendice: hipoteses
sl = novo_slide(None, 'Apêndice — hipóteses declaradas', 'Para a arguição: o que é dado de catálogo e o que é escolha da equipe')
hip = [('H1', 'fase do enrolamento: ¼ de passo', 'evita empate numérico; amplitude do erro não muda'),
       ('H2', 'folga calibrada na subida', 'mesma subida usada na reta terminal da E01'),
       ('H3', 'ruído 1/f, média nula, pico = ENR', 'forma 1/f é a atribuição; o datasheet só dá o pico'),
       ('H4', 'I_cursor = I_B máx. do OPA333 (200 pA)', 'pior caso de catálogo a 25 °C'),
       ('H5', 'ENR só com o cursor em movimento', 'definição do fabricante'),
       ('H6', '“histerese/contato” = folga eixo–cursor', 'única histerese especificada no 3590'),
       ('H7', 'contato de base 110 Ω (só no contrafactual)', 'para o resistor variável nunca zerar')]
yy = Cm(3.6)
for cod, h, pq in hip:
    texto(sl, Cm(1.2), yy, Cm(1.5), Cm(1.0), cod, tam=15, negrito=True, cor=AZUL)
    texto(sl, Cm(2.8), yy, Cm(13.5), Cm(1.0), h, tam=15)
    texto(sl, Cm(16.6), yy, Cm(16), Cm(1.0), pq, tam=14, cor=TINTA2)
    yy += Cm(1.75)

os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
prs.save(SAIDA)
print('apresentacao salva em', os.path.abspath(SAIDA), '| slides:', n_slide[0],
      '| TyphoonSim:', 'sim' if TY else 'pendente')
